import io
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

def generate_pdf_report(report_obj) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontSize=18, textColor=colors.HexColor('#1E293B'))
    heading_style = ParagraphStyle('HeadingStyle', parent=styles['Heading2'], fontSize=14, textColor=colors.HexColor('#0F172A'))
    normal_style = styles['Normal']

    story.append(Paragraph("Legal Metrology Compliance Report", title_style))
    story.append(Spacer(1, 10))

    prod_title = report_obj.product.title if report_obj.product else "N/A"
    prod_url = report_obj.product.url if report_obj.product else "N/A"
    cat_name = report_obj.category.name if report_obj.category else "Uncategorized"

    # Meta Table
    meta_data = [
        ["Report ID:", str(report_obj.id), "Date:", report_obj.created_at.strftime("%Y-%m-%d %H:%M")],
        ["Product Title:", prod_title[:40], "Category:", cat_name],
        ["Product URL:", prod_url[:50], "Overall Status:", report_obj.overall_status],
        ["Compliance Score:", f"{report_obj.overall_score}%", "OCR Checked:", "Yes" if report_obj.ocr_checked else "No"]
    ]
    t = Table(meta_data, colWidths=[100, 170, 90, 170])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('FONTNAME', (0,0), (-1,-1), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 9),
    ]))
    story.append(t)
    story.append(Spacer(1, 15))

    # Violations & Rules Summary Table
    story.append(Paragraph("Detailed Rules Evaluation", heading_style))
    story.append(Spacer(1, 8))

    rules_table_data = [["Rule Code", "Declaration", "Status", "Actual Found Value", "AI Rationale"]]
    for v in report_obj.violations:
        status_color = "#16A34A" if v.status == "PASSED" else "#DC2626"
        rule_code = v.rule.rule_code if v.rule else "RULE"
        decl_name = v.rule.declaration_name if v.rule else v.field_name
        rules_table_data.append([
            rule_code,
            decl_name,
            v.status,
            (v.actual_found or "Not Found")[:30],
            (v.ai_reasoning or "")[:40]
        ])

    rt = Table(rules_table_data, colWidths=[60, 110, 70, 130, 160])
    rt.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0F172A')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('FONTSIZE', (0,0), (-1,-1), 8),
    ]))
    story.append(rt)

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()

def generate_excel_report(report_obj) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Compliance Report"

    ws.append(["Report ID", report_obj.id])
    ws.append(["Product URL", report_obj.product.url if report_obj.product else ""])
    ws.append(["Product Title", report_obj.product.title if report_obj.product else ""])
    ws.append(["Detected Category", report_obj.category.name if report_obj.category else ""])
    ws.append(["Overall Compliance Status", report_obj.overall_status])
    ws.append(["Overall Score", f"{report_obj.overall_score}%"])
    ws.append([])

    ws.append(["Rule Code", "Declaration Name", "Enforcement Level", "Status", "Actual Value Found", "AI Rationale"])
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col in range(1, 7):
        cell = ws.cell(row=8, column=col)
        cell.fill = header_fill
        cell.font = header_font

    row_idx = 9
    for v in report_obj.violations:
        ws.append([
            v.rule.rule_code if v.rule else "RULE",
            v.rule.declaration_name if v.rule else v.field_name,
            v.rule.enforcement_level if v.rule else "strict",
            v.status,
            v.actual_found or "Not Found",
            v.ai_reasoning or ""
        ])
        row_idx += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
